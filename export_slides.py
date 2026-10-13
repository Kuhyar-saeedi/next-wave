#!/usr/bin/env python3
"""Export every deck as a PDF and a PowerPoint file, slides only (no speaker notes).

Slides are rendered in Chromium exactly as the site shows them, so their
content and look stay unchanged:
  - PDF: one 1920x1080 page per slide, with real (selectable) text.
  - PPTX: one 16:9 slide per slide, holding a full-size picture of it.
Files are written to src/downloads/ and published by build.py.

Run after changing any slides:   python3 export_slides.py
Requires: playwright (with Chromium), python-pptx, Pillow.
"""
import functools
import http.server
import io
import re
import shutil
import threading
from pathlib import Path

from PIL import Image
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

import build

ROOT = Path(__file__).parent
DL = build.SRC / "downloads"
TMP = build.OUT / "_export"
PORT = 8799

FONT_CSS = """
@font-face { font-family: 'DM Sans'; font-style: normal; font-weight: 100 1000; font-display: block;
  src: url(../assets/fonts/dm-sans-latin-wght-normal.woff2) format('woff2-variations');
  unicode-range: U+0000-00FF, U+0131, U+0152-0153, U+02BB-02BC, U+02C6, U+02DA, U+02DC, U+0304, U+0308, U+0329, U+2000-206F, U+20AC, U+2122, U+2191, U+2193, U+2212, U+2215, U+FEFF, U+FFFD; }
@font-face { font-family: 'DM Sans'; font-style: normal; font-weight: 100 1000; font-display: block;
  src: url(../assets/fonts/dm-sans-latin-ext-wght-normal.woff2) format('woff2-variations');
  unicode-range: U+0100-02BA, U+02BD-02C5, U+02C7-02CC, U+02CE-02D7, U+02DD-02FF, U+0304, U+0308, U+0329, U+1D00-1DBF, U+1E00-1E9F, U+1EF2-1EFF, U+2020, U+20A0-20AB, U+20AD-20C0, U+2113, U+2C60-2C7F, U+A720-A7FF; }
"""

PAGE_CSS = """
@page { size: 1920px 1080px; margin: 0; }
html, body { margin: 0; padding: 0; background: #fff; }
.deck { display: block !important; gap: 0 !important; }
.slide-frame { width: 1920px !important; height: 1080px !important; aspect-ratio: auto !important;
  border-radius: 0 !important; box-shadow: none !important; overflow: hidden; break-after: page; page-break-after: always; }
.slide-frame:last-child { break-after: auto; page-break-after: auto; }
"""


def slug(text):
    s = re.sub(r"[—–]", "-", text)
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")
    return re.sub(r"-+", "-", s)


def file_stem(key):
    return "Next-Wave-" + slug(build.deck_title(key))


def export_page(key):
    frames = []
    for sid in build.decks[key]["order"]:
        own = build.DECK_DIR / key / "slides" / f"{sid}.html"
        raw = (own if own.exists() else build.slide_files[sid]).read_text()
        sec = re.sub(r"\s*<aside>.*?</aside>", "", raw, flags=re.S).strip()  # drop speaker notes
        frames.append(f'<div class="slide-frame">{sec}</div>')
    return f"""<!doctype html><html><head><meta charset="utf-8">
<link rel="stylesheet" href="../assets/style.css">
<style>{FONT_CSS}{PAGE_CSS}</style></head>
<body><div class="deck">{''.join(frames)}</div>
<script src="../assets/slides.js"></script></body></html>"""


def serve():
    class Quiet(http.server.SimpleHTTPRequestHandler):
        def log_message(self, *a):
            pass
    handler = functools.partial(Quiet, directory=str(build.OUT))
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", PORT), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def make_pptx(images, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in, 16:9
    blank = prs.slide_layouts[6]
    for png in images:
        im = Image.open(io.BytesIO(png)).convert("RGB")
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=92, optimize=True)
        buf.seek(0)
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(path)


def main():
    build.main()
    TMP.mkdir(exist_ok=True)
    if DL.exists():
        shutil.rmtree(DL)
    DL.mkdir(parents=True)
    srv = serve()
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1920, "height": 1080})
            for key in build.decks:
                (TMP / f"{key}.html").write_text(export_page(key))
                page.goto(f"http://127.0.0.1:{PORT}/_export/{key}.html", wait_until="networkidle")
                page.evaluate("document.fonts.ready")
                page.wait_for_timeout(300)
                stem = file_stem(key)
                page.pdf(path=str(DL / f"{stem}.pdf"), width="1920px", height="1080px",
                         print_background=True, prefer_css_page_size=True)
                frames = page.locator(".slide-frame")
                shots = [frames.nth(i).screenshot(type="png") for i in range(frames.count())]
                make_pptx(shots, DL / f"{stem}.pptx")
                print(f"{key}: {len(shots)} slides -> {stem}.pdf / .pptx")
            browser.close()
    finally:
        srv.shutdown()
        shutil.rmtree(TMP, ignore_errors=True)
    build.main()  # rebuild so the site picks up the new files


if __name__ == "__main__":
    main()
